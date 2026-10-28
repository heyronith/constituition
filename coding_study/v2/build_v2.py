"""Build simplified (v2) human-coding packages for two new coders (C, D).

Stimulus source: coder A's original package sheet (same 120 rows as the Phase 12 key).
Outputs (local): out/coderC, out/coderD, out/common, mapping_v2.csv, zips.
Usage: python build_v2.py <coderA_sheet.xlsx> <outdir>
"""
import sys, os, json, random, csv, datetime, subprocess, zipfile, shutil, hashlib
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261004
CODERS = {"C": SEED + 11, "D": SEED + 12}
HERE = os.path.dirname(os.path.abspath(__file__))
MAT = os.path.join(HERE, "materials")
FIXED_TIME = datetime.datetime(2026, 10, 27, 0, 0, 0)

HEAD = ["row", "ORIGINAL", "REVISED", "Q1_demand", "Q2_decides_itself", "Q3_overall", "notes"]
LEGEND = [
    ("Q1_demand", "Weaker", "The revised rule asks for LESS (softer wording, any new exception, fewer situations, part removed, or now allows what it forbade)."),
    ("Q1_demand", "Same", "Different words, same demand. Also: the only change is 'authorised/verified people only' (note 'approved people')."),
    ("Q1_demand", "Stronger", "The revised rule asks for MORE."),
    ("Q1_demand", "Reminder", "Q1 is about how much the rule demands, NOT whether the change is good (that is Q3)."),
    ("Q2_decides_itself", "Yes", "Only if Q1 = Weaker: the assistant itself decides when the rule applies, or may refuse based on its own judgment, beliefs, values or principles."),
    ("Q2_decides_itself", "No", "Only if Q1 = Weaker: the exception depends on something outside the assistant, or the rule was simply softened or narrowed."),
    ("Q3_overall", "Worse / Neutral / Better / Unsure", "Your personal opinion: is the change worse, neutral or better for the people the assistant serves and society?"),
]


def load_stimuli(path):
    ws = openpyxl.load_workbook(path)["Coding"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 120
    return [(r[0], r[1], r[2]) for r in rows]  # (coderA_row, ORIGINAL, REVISED)


def make_sheet(path, items, with_timelog):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Coding"
    ws.append(HEAD)
    for i, (o, r) in enumerate(items, 1):
        ws.append([i, o, r, None, None, None, None])
    n = len(items) + 1
    widths = {"A": 6, "B": 60, "C": 60, "D": 14, "E": 18, "F": 14, "G": 40}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    shade = PatternFill("solid", fgColor="F2F2F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=1, max_row=n):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column_letter in "ABC" and cell.row > 1:
                cell.fill = shade
            if cell.column_letter in "DEFG" and cell.row > 1:
                cell.protection = openpyxl.styles.Protection(locked=False)
    ws.freeze_panes = "B2"
    for col, opts in (("D", "Weaker,Same,Stronger"), ("E", "Yes,No"), ("F", "Worse,Neutral,Better,Unsure")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True, showDropDown=False)
        dv.error = "Please choose from the list."
        dv.errorTitle = "Invalid entry"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    ws.protection.sheet = True
    lg = wb.create_sheet("Legend")
    lg.append(["Column", "Answer", "Meaning"])
    for r in LEGEND:
        lg.append(list(r))
    lg.column_dimensions["A"].width = 20
    lg.column_dimensions["B"].width = 26
    lg.column_dimensions["C"].width = 100
    for row in lg.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in lg[1]:
        cell.font = Font(bold=True)
    if with_timelog:
        tl = wb.create_sheet("Time log")
        tl.append(["Sitting", "Date", "Start time", "End time", "Rows done (e.g. 1-40)"])
        for k in range(1, 9):
            tl.append([k, None, None, None, None])
        for c in "ABCDE":
            tl.column_dimensions[c].width = 22
        for cell in tl[1]:
            cell.font = Font(bold=True)
    wb.properties.creator = "coding_study"
    wb.properties.lastModifiedBy = "coding_study"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.save(path)


def pandoc_pdf(md_path, pdf_path, landscape=False, fontsize="11pt"):
    if os.environ.get("SKIP_PDF"):
        open(pdf_path, "wb").close(); return
    cmd = ["pandoc", md_path, "-s", "-o", pdf_path, "--pdf-engine=wkhtmltopdf",
           "--css", os.path.join(MAT, "style.css"), "-V", "papersize=letter",
           "-V", "margin-top=14mm", "-V", "margin-bottom=14mm", "-V", "margin-left=12mm", "-V", "margin-right=12mm",
           "--pdf-engine-opt=--quiet", "--pdf-engine-opt=--disable-smart-shrinking"]
    if landscape:
        cmd += ["--pdf-engine-opt=--orientation", "--pdf-engine-opt=Landscape"]
    if fontsize != "11pt":
        cmd += ["--metadata", "pagetitle=x", "-V", f"header-includes=<style>body{{font-size:{fontsize}}}</style>"]
    subprocess.run(cmd, check=True)


def main(src, out):
    os.makedirs(out, exist_ok=True)
    common = os.path.join(out, "common"); os.makedirs(common, exist_ok=True)
    stim = load_stimuli(src)
    practice = json.load(open(os.path.join(MAT, "practice_items_v2.json")))

    # PDFs
    pandoc_pdf(os.path.join(MAT, "CODING_GUIDE_v2.md"), os.path.join(common, "Coding_Guide.pdf"))
    ans = ["---", 'title: "Practice answers"', "---", "",
           "All items are invented. Q3 answers are only suggestions: Q3 is a personal opinion, and differences there are expected.", "",
           "| # | ORIGINAL | REVISED | Q1 | Q2 | Q3 (suggested) | Why |",
           "|--:|:-----------------|:-----------------|:--|:--|:--|:-----------------------|"]
    for p in practice:
        ans.append(f"| {p['n']} | {p['original']} | {p['revised']} | {p['q1']} | {p['q2'] or '—'} | {p['q3']} | {p['why']} |")
    ans_md = os.path.join(common, "Practice_Answers.md"); open(ans_md, "w").write("\n".join(ans) + "\n")
    pandoc_pdf(ans_md, os.path.join(common, "Practice_Answers.pdf"), landscape=True, fontsize="9pt")
    start = ["---", 'title: "Start here"', "---", "",
             "1. Read **Coding_Guide.pdf** (about 10 minutes).",
             "2. Fill in **practice_sheet.xlsx** (10 rows) **without looking anything up**, and send it back to the person who sent you this package.",
             "3. You will then receive the **answer key** and your **main sheet** (120 rows).", "",
             "Please **work alone**, **do not use AI tools** (ChatGPT, Copilot, Gemini, etc.), and write your **start and end times** in the \"Time log\" tab of the main sheet.", "",
             "Thank you!"]
    st_md = os.path.join(common, "Start_Here.md"); open(st_md, "w").write("\n".join(start) + "\n")
    pandoc_pdf(st_md, os.path.join(common, "Start_Here.pdf"))
    make_sheet(os.path.join(common, "practice_sheet.xlsx"), [(p["original"], p["revised"]) for p in practice], with_timelog=False)

    # coder sheets
    mapping = []
    for c, seed in CODERS.items():
        order = list(range(120)); random.Random(seed).shuffle(order)
        items = [(stim[i][1], stim[i][2]) for i in order]
        d = os.path.join(out, f"coder{c}"); os.makedirs(d, exist_ok=True)
        make_sheet(os.path.join(d, f"coding_sheet_coder{c}.xlsx"), items, with_timelog=True)
        for row, i in enumerate(order, 1):
            mapping.append({"coder": c, "row": row, "coderA_row": stim[i][0]})
        # zips
        z1 = os.path.join(out, f"package_coder{c}_STEP1.zip")
        with zipfile.ZipFile(z1, "w", zipfile.ZIP_DEFLATED) as z:
            for f in ("Start_Here.pdf", "Coding_Guide.pdf", "practice_sheet.xlsx"):
                z.write(os.path.join(common, f), f)
        z2 = os.path.join(out, f"package_coder{c}_STEP2.zip")
        with zipfile.ZipFile(z2, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(os.path.join(common, "Practice_Answers.pdf"), "Practice_Answers.pdf")
            z.write(os.path.join(d, f"coding_sheet_coder{c}.xlsx"), f"coding_sheet_coder{c}.xlsx")
    with open(os.path.join(out, "mapping_v2.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["coder", "row", "coderA_row"]); w.writeheader(); w.writerows(mapping)
    print("built", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])

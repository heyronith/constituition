#!/usr/bin/env python3
"""Analyze returned human-coding sheets (EXPLORATORY)."""

from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261004
LABEL = "EXPLORATORY (not preregistered)"
Q1_ALLOWED = frozenset(list("ABCDEF"))
Q3_ALLOWED = frozenset(list("WNBU"))
THEME_COLS = [f"T{i}" for i in range(1, 8)]
EROSION_CODES = frozenset(list("ABC"))

FATE_TO_Q1 = {
    "INVERTED": "A",
    "SUBORDINATED": "B",
    "WEAKENED": "C",
    "QUALIFIED_LEGITIMACY": "D",
    "STRENGTHENED": "E",
    "RETAINED": "F",
}


def study_root() -> Path:
    return Path(__file__).resolve().parent


def _norm_ws(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())


def dump_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")


def read_coding_sheet(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, data_only=True)
    ws = wb["Coding"] if "Coding" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    data = []
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        data.append({header[i]: (r[i] if i < len(r) else None) for i in range(len(header))})
    df = pd.DataFrame(data)
    # normalize
    if "row" in df.columns:
        df["row"] = df["row"].astype(int)
    def _clean_code(v: Any) -> str | None:
        if v is None:
            return None
        if isinstance(v, float) and math.isnan(v):
            return None
        s = str(v).strip()
        if s == "" or s.lower() == "nan":
            return None
        return s.upper()

    for c in ["Q1_change", "Q3_overall", *THEME_COLS]:
        if c in df.columns:
            df[c] = df[c].map(_clean_code)
    return df


def _as_code(v: Any) -> str | None:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    s = str(v).strip()
    if s == "" or s.lower() == "nan":
        return None
    return s.upper()


def validate_sheet(
    df: pd.DataFrame, key: pd.DataFrame, sheet_map: pd.DataFrame, coder: str
) -> list[str]:
    problems: list[str] = []
    m = sheet_map[sheet_map.coder == coder].copy()
    if len(m) != len(df):
        problems.append(f"coder {coder}: expected {len(m)} rows, got {len(df)}")
    merged = df.merge(m, on="row", how="left").merge(key, on="key_id", how="left", suffixes=("", "_key"))
    for _, r in merged.iterrows():
        row = int(r["row"])
        q1 = _as_code(r.get("Q1_change"))
        q3 = _as_code(r.get("Q3_overall"))
        if q1 not in Q1_ALLOWED:
            problems.append(f"coder {coder} row {row}: missing/invalid Q1={q1!r}")
        if q3 not in Q3_ALLOWED:
            problems.append(f"coder {coder} row {row}: missing/invalid Q3={q3!r}")
        themes = [_as_code(r.get(t)) for t in THEME_COLS]
        for t, v in zip(THEME_COLS, themes):
            if v not in (None, "Y"):
                problems.append(f"coder {coder} row {row}: invalid {t}={v!r}")
        if q1 not in EROSION_CODES and any(v == "Y" for v in themes):
            problems.append(f"coder {coder} row {row}: themes set but Q1={q1} not in A/B/C")
        # text unchanged
        if _norm_ws(str(r.get("ORIGINAL") or "")) != _norm_ws(str(r.get("original") or "")):
            problems.append(f"coder {coder} row {row}: ORIGINAL text altered")
        if _norm_ws(str(r.get("REVISED") or "")) != _norm_ws(str(r.get("revised") or "")):
            problems.append(f"coder {coder} row {row}: REVISED text altered")
    return problems


def cohen_kappa(y1: list[Any], y2: list[Any]) -> float:
    if len(y1) != len(y2) or not y1:
        return float("nan")
    labels = sorted(set(y1) | set(y2), key=lambda x: str(x))
    idx = {lab: i for i, lab in enumerate(labels)}
    n = len(y1)
    k = len(labels)
    mat = np.zeros((k, k), dtype=float)
    for a, b in zip(y1, y2):
        mat[idx[a], idx[b]] += 1
    po = float(np.trace(mat) / n)
    pe = float(sum(mat.sum(0)[i] * mat.sum(1)[i] for i in range(k)) / (n * n))
    if pe == 1.0:
        return 1.0 if po == 1.0 else 0.0
    return (po - pe) / (1.0 - pe)


def krippendorff_alpha_nominal(c1: list[Any], c2: list[Any]) -> float:
    """Nominal α for two coders, one rating each per unit."""
    if len(c1) != len(c2) or not c1:
        return float("nan")
    pairs = list(zip(c1, c2))
    # coincidence matrix
    values = sorted(set(c1) | set(c2), key=lambda x: str(x))
    idx = {v: i for i, v in enumerate(values)}
    m = len(values)
    o = np.zeros((m, m), dtype=float)
    for a, b in pairs:
        i, j = idx[a], idx[b]
        o[i, j] += 1
        o[j, i] += 1
        # each pair contributes 2 to coincidence totals in Krippendorff's formulation
    # Better standard 2-coder formula:
    # Do = proportion of disagreeing pairs
    n = len(pairs)
    disagree = sum(1 for a, b in pairs if a != b)
    do = disagree / n
    # expected
    counts = Counter(list(c1) + list(c2))
    n_rat = 2 * n
    de = 0.0
    for v, nv in counts.items():
        for w, nw in counts.items():
            if v == w:
                continue
            de += (nv / n_rat) * ((nw) / (n_rat - 1))
    # For nominal, De = 1 - sum p_v^2 adjusted; use coincidence form:
    # Simpler: de = 1 - sum_c n_c(n_c-1) / (N(N-1)) where coincidence within
    n_vals = np.array([counts[v] for v in values], dtype=float)
    de = 1.0 - float(np.sum(n_vals * (n_vals - 1)) / (n_rat * (n_rat - 1))) if n_rat > 1 else 0.0
    if de == 0.0:
        return 1.0 if do == 0.0 else 0.0
    return 1.0 - (do / de)


def bootstrap_ci(
    values_fn, n: int, rng: np.random.Generator, n_boot: int = 1000
) -> tuple[float, float, float]:
    point = values_fn(np.arange(n))
    boots = []
    for _ in range(n_boot):
        idx = rng.integers(0, n, size=n)
        boots.append(values_fn(idx))
    lo, hi = np.nanquantile(boots, [0.025, 0.975])
    return float(point), float(lo), float(hi)


def confusion(y_true: list[str], y_pred: list[str], labels: list[str]) -> dict[str, Any]:
    mat = {a: {b: 0 for b in labels} for a in labels}
    for t, p in zip(y_true, y_pred):
        if t not in mat:
            mat[t] = {b: 0 for b in labels}
        if p not in mat[t]:
            for row in mat.values():
                row.setdefault(p, 0)
            mat[t][p] = 0
        mat[t][p] = mat[t].get(p, 0) + 1
    agree = sum(1 for t, p in zip(y_true, y_pred) if t == p)
    return {
        "labels": labels,
        "matrix": mat,
        "n": len(y_true),
        "pct_agreement": (agree / len(y_true)) if y_true else None,
        "cohen_kappa": cohen_kappa(y_true, y_pred) if y_true else None,
    }


def analyze(
    sheet_a: Path,
    sheet_b: Path,
    *,
    key_path: Path,
    map_path: Path,
    out_dir: Path,
) -> dict[str, Any]:
    key = pd.read_csv(key_path)
    sheet_map = pd.read_csv(map_path)
    sheet_map["coder"] = sheet_map["coder"].astype(str)
    a = read_coding_sheet(sheet_a)
    b = read_coding_sheet(sheet_b)

    problems = validate_sheet(a, key, sheet_map, "A") + validate_sheet(b, key, sheet_map, "B")
    if problems:
        print("VALIDATION FAILED:")
        for p in problems:
            print(" -", p)
        raise SystemExit(1)

    def join_coder(df: pd.DataFrame, coder: str) -> pd.DataFrame:
        m = sheet_map[sheet_map.coder == coder][["row", "key_id"]]
        return df.merge(m, on="row").merge(key, on="key_id", suffixes=("", "_key"))

    ja = join_coder(a, "A").sort_values("key_id").reset_index(drop=True)
    jb = join_coder(b, "B").sort_values("key_id").reset_index(drop=True)
    assert list(ja.key_id) == list(jb.key_id)

    # anonymised codes
    codes = pd.DataFrame(
        {
            "key_id": ja.key_id,
            "stratum": ja.stratum,
            "item_id": ja.item_id,
            "category": ja.category,
            "fate_gpt54": ja.fate_gpt54,
            "fate_mimo": ja.fate_mimo,
            "coder_A_Q1": ja.Q1_change,
            "coder_B_Q1": jb.Q1_change,
            "coder_A_Q3": ja.Q3_overall,
            "coder_B_Q3": jb.Q3_overall,
            **{f"coder_A_{t}": ja[t] for t in THEME_COLS},
            **{f"coder_B_{t}": jb[t] for t in THEME_COLS},
            "original": ja.original,
            "revised": ja.revised,
        }
    )
    out_dir.mkdir(parents=True, exist_ok=True)
    codes.to_csv(out_dir / "codes_anonymised.csv", index=False)

    rng = np.random.default_rng(SEED)

    # Reliability
    q1_a = ja.Q1_change.tolist()
    q1_b = jb.Q1_change.tolist()
    er_a = [1 if x in EROSION_CODES else 0 for x in q1_a]
    er_b = [1 if x in EROSION_CODES else 0 for x in q1_b]

    def alpha_q1(idx):
        return krippendorff_alpha_nominal([q1_a[i] for i in idx], [q1_b[i] for i in idx])

    def alpha_bin(idx):
        return krippendorff_alpha_nominal([er_a[i] for i in idx], [er_b[i] for i in idx])

    a_q1, a_q1_lo, a_q1_hi = bootstrap_ci(alpha_q1, len(q1_a), rng)
    a_bin, a_bin_lo, a_bin_hi = bootstrap_ci(alpha_bin, len(er_a), rng)

    # Themes κ on rows where both A/B/C
    both_erosion = [i for i, (x, y) in enumerate(zip(q1_a, q1_b)) if x in EROSION_CODES and y in EROSION_CODES]
    theme_kappa = {}
    for t in THEME_COLS:
        ya = [(ja.iloc[i][t] == "Y") for i in both_erosion]
        yb = [(jb.iloc[i][t] == "Y") for i in both_erosion]
        theme_kappa[t] = {
            "n": len(both_erosion),
            "cohen_kappa": cohen_kappa(ya, yb) if both_erosion else None,
        }

    # Q3 α excluding U
    q3_pairs = [
        (i, ja.iloc[i].Q3_overall, jb.iloc[i].Q3_overall)
        for i in range(len(ja))
        if ja.iloc[i].Q3_overall != "U" and jb.iloc[i].Q3_overall != "U"
    ]
    if q3_pairs:
        idxs = [p[0] for p in q3_pairs]
        q3a = [p[1] for p in q3_pairs]
        q3b = [p[2] for p in q3_pairs]

        def alpha_q3(boot_idx):
            # boot_idx indexes into q3_pairs
            return krippendorff_alpha_nominal([q3a[i] for i in boot_idx], [q3b[i] for i in boot_idx])

        a_q3, a_q3_lo, a_q3_hi = bootstrap_ci(alpha_q3, len(q3_pairs), rng)
    else:
        a_q3 = a_q3_lo = a_q3_hi = float("nan")

    reliability = {
        "label": LABEL,
        "krippendorff_alpha_Q1_nominal": {
            "alpha": a_q1,
            "ci95": [a_q1_lo, a_q1_hi],
            "n": len(q1_a),
        },
        "krippendorff_alpha_binary_erosion": {
            "alpha": a_bin,
            "ci95": [a_bin_lo, a_bin_hi],
            "n": len(er_a),
        },
        "cohen_kappa_themes_both_ABC": theme_kappa,
        "krippendorff_alpha_Q3_excl_U": {
            "alpha": a_q3,
            "ci95": [a_q3_lo, a_q3_hi],
            "n": len(q3_pairs),
        },
    }
    dump_json(out_dir / "reliability.json", reliability)

    # Human vs judge
    labels6 = list("ABCDEF")
    judge_q1 = [FATE_TO_Q1.get(str(f), None) for f in ja.fate_gpt54.tolist()]
    # drop rows where fate not in map (shouldn't happen for sampled)
    hvj: dict[str, Any] = {"label": LABEL, "gpt54": {}, "mimo": {}}

    def vs_judge(coder_q1: list[str], judge: list[str | None], name: str) -> dict[str, Any]:
        pairs = [(c, j) for c, j in zip(coder_q1, judge) if j is not None]
        if not pairs:
            return {"n": 0}
        c_list = [p[0] for p in pairs]
        j_list = [p[1] for p in pairs]
        six = confusion(j_list, c_list, labels6)
        j_bin = [1 if j in EROSION_CODES else 0 for j in j_list]
        c_bin = [1 if c in EROSION_CODES else 0 for c in c_list]
        bin_ = confusion([str(x) for x in j_bin], [str(x) for x in c_bin], ["0", "1"])
        return {"six_way": six, "binary_erosion": bin_}

    hvj["gpt54"]["coder_A"] = vs_judge(q1_a, judge_q1, "A")
    hvj["gpt54"]["coder_B"] = vs_judge(q1_b, judge_q1, "B")
    consensus_q1 = [a if a == b else None for a, b in zip(q1_a, q1_b)]
    cons_pairs_q1 = [c for c in consensus_q1 if c is not None]
    cons_judge = [j for c, j in zip(consensus_q1, judge_q1) if c is not None]
    hvj["gpt54"]["consensus"] = vs_judge(cons_pairs_q1, cons_judge, "consensus")

    mimo_q1 = [FATE_TO_Q1.get(str(f), None) if pd.notna(f) else None for f in ja.fate_mimo.tolist()]
    hvj["mimo"]["coder_A"] = vs_judge(q1_a, mimo_q1, "A")
    hvj["mimo"]["coder_B"] = vs_judge(q1_b, mimo_q1, "B")
    cons_mimo = [j for c, j in zip(consensus_q1, mimo_q1) if c is not None]
    hvj["mimo"]["consensus"] = vs_judge(cons_pairs_q1, cons_mimo, "consensus")
    dump_json(out_dir / "human_vs_judge.json", hvj)

    # Themes by stratum
    theme_rows = []
    q3_rows = []
    for stratum, g_idx in codes.groupby("stratum").groups.items():
        idxs = list(g_idx)
        n = len(idxs)
        for t in THEME_COLS:
            both = sum(
                1
                for i in idxs
                if codes.loc[i, f"coder_A_{t}"] == "Y" and codes.loc[i, f"coder_B_{t}"] == "Y"
            )
            either = sum(
                1
                for i in idxs
                if codes.loc[i, f"coder_A_{t}"] == "Y" or codes.loc[i, f"coder_B_{t}"] == "Y"
            )
            theme_rows.append(
                {
                    "label": LABEL,
                    "stratum": stratum,
                    "theme": t,
                    "n_rows": n,
                    "n_both": both,
                    "pct_both": both / n if n else None,
                    "n_either": either,
                    "pct_either": either / n if n else None,
                }
            )
        # Q3 distribution: count each coder separately + agree
        for coder, col in [("A", "coder_A_Q3"), ("B", "coder_B_Q3")]:
            ctr = Counter(codes.loc[i, col] for i in idxs)
            for k, v in ctr.items():
                q3_rows.append(
                    {
                        "label": LABEL,
                        "stratum": stratum,
                        "coder": coder,
                        "Q3": k,
                        "n": int(v),
                        "pct": v / n if n else None,
                    }
                )
    pd.DataFrame(theme_rows).to_csv(out_dir / "themes_by_stratum.csv", index=False)
    pd.DataFrame(q3_rows).to_csv(out_dir / "q3_by_stratum.csv", index=False)

    # Disagreements
    disag = []
    for i in range(len(ja)):
        if ja.iloc[i].Q1_change != jb.iloc[i].Q1_change or ja.iloc[i].Q3_overall != jb.iloc[i].Q3_overall:
            disag.append(
                {
                    "key_id": ja.iloc[i].key_id,
                    "stratum": ja.iloc[i].stratum,
                    "coder_A_Q1": ja.iloc[i].Q1_change,
                    "coder_B_Q1": jb.iloc[i].Q1_change,
                    "coder_A_Q3": ja.iloc[i].Q3_overall,
                    "coder_B_Q3": jb.iloc[i].Q3_overall,
                    "original": ja.iloc[i].original,
                    "revised": ja.iloc[i].revised,
                }
            )
    pd.DataFrame(disag).to_csv(out_dir / "disagreements.csv", index=False)

    # summary
    lines = [
        f"# Human coding analysis — {LABEL}",
        "",
        f"- n keys: {len(ja)}",
        f"- Krippendorff α Q1: {a_q1:.3f} [{a_q1_lo:.3f}, {a_q1_hi:.3f}]",
        f"- Krippendorff α binary erosion: {a_bin:.3f} [{a_bin_lo:.3f}, {a_bin_hi:.3f}]",
        f"- Krippendorff α Q3 (excl U): {a_q3:.3f} [{a_q3_lo:.3f}, {a_q3_hi:.3f}]",
        f"- disagreements (Q1 or Q3): {len(disag)}",
        f"- GPT-5.4 vs coder A binary pct agree: {hvj['gpt54']['coder_A']['binary_erosion']['pct_agreement']}",
        f"- GPT-5.4 vs coder B binary pct agree: {hvj['gpt54']['coder_B']['binary_erosion']['pct_agreement']}",
        "",
    ]
    (out_dir / "summary.md").write_text("\n".join(lines), encoding="utf-8")
    return {"label": LABEL, "n": len(ja), "n_disagreements": len(disag), "reliability": reliability}


# ---------------------------------------------------------------------------
# Synthetic dry-run helpers
# ---------------------------------------------------------------------------

def _write_filled_sheet(
    path: Path,
    key: pd.DataFrame,
    sheet_map: pd.DataFrame,
    coder: str,
    seed: int,
) -> None:
    import importlib.util

    bp_path = study_root() / "build_package.py"
    spec = importlib.util.spec_from_file_location("coding_build_package", bp_path)
    assert spec and spec.loader
    bp = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(bp)
    SHEET_COLUMNS = bp.SHEET_COLUMNS
    _add_legend = bp._add_legend
    _style_sheet = bp._style_sheet

    rng = np.random.default_rng(seed)
    m = sheet_map[sheet_map.coder == coder].sort_values("row")
    merged = m.merge(key, on="key_id")
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Coding"
    ws.append(SHEET_COLUMNS)
    q1_choices = list("ABCDEF")
    q3_choices = list("WNBU")
    for _, r in merged.iterrows():
        q1 = q1_choices[int(rng.integers(0, len(q1_choices)))]
        themes = [""] * 7
        if q1 in EROSION_CODES:
            for j in range(7):
                if rng.random() < 0.25:
                    themes[j] = "Y"
        q3 = q3_choices[int(rng.integers(0, len(q3_choices)))]
        ws.append([int(r["row"]), r["original"], r["revised"], q1, *themes, q3, ""])
    _style_sheet(ws, len(merged))
    _add_legend(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def synthetic_dry_run(repo_out: Path | None = None) -> dict[str, Any]:
    base = study_root()
    out = repo_out or (base / "out")
    key = pd.read_csv(out / "sampling_key.csv")
    sheet_map = pd.read_csv(out / "sheet_map.csv")
    syn = out / "synthetic"
    syn.mkdir(parents=True, exist_ok=True)
    path_a = syn / "coding_sheet_coderA_synthetic.xlsx"
    path_b = syn / "coding_sheet_coderB_synthetic.xlsx"
    _write_filled_sheet(path_a, key, sheet_map, "A", SEED + 3)
    _write_filled_sheet(path_b, key, sheet_map, "B", SEED + 4)
    results = out / "synthetic_results"
    if results.exists():
        import shutil

        shutil.rmtree(results)
    return analyze(path_a, path_b, key_path=out / "sampling_key.csv", map_path=out / "sheet_map.csv", out_dir=results)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("sheet_a", type=Path, nargs="?", help="Returned coder A xlsx")
    p.add_argument("sheet_b", type=Path, nargs="?", help="Returned coder B xlsx")
    p.add_argument("--key", type=Path, default=None)
    p.add_argument("--map", type=Path, default=None)
    p.add_argument("--out", type=Path, default=None)
    p.add_argument("--synthetic-dry-run", action="store_true")
    args = p.parse_args()
    base = study_root()
    if args.synthetic_dry_run:
        info = synthetic_dry_run()
        print(json.dumps({"dry_run": True, "n": info["n"], "n_disagreements": info["n_disagreements"]}, indent=2))
        return
    if not args.sheet_a or not args.sheet_b:
        p.error("sheet_a and sheet_b required unless --synthetic-dry-run")
    key = args.key or (base / "out" / "sampling_key.csv")
    mp = args.map or (base / "out" / "sheet_map.csv")
    out = args.out or (repo_root_results())
    info = analyze(args.sheet_a, args.sheet_b, key_path=key, map_path=mp, out_dir=out)
    print(json.dumps({"n": info["n"], "n_disagreements": info["n_disagreements"]}, indent=2))


def repo_root_results() -> Path:
    from rc.config import repo_root

    return repo_root() / "results" / "exploratory_aie" / "human_coding"


if __name__ == "__main__":
    main()

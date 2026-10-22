"""Tests for the blind human-coding package."""

from __future__ import annotations

import filecmp
import importlib.util
import shutil
import zipfile
from pathlib import Path

import pandas as pd
import pytest

from rc.config import repo_root

ROOT = repo_root()
CODING = ROOT / "runs" / "main_v1" / "coding"
STUDY = ROOT / "coding_study"
OUT = STUDY / "out"
BUILD = STUDY / "build_package.py"
ANALYZE = STUDY / "analyze_coding.py"

pytestmark = pytest.mark.skipif(not CODING.exists(), reason="runs/main_v1/coding/ missing")


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def built_package():
    build = _load(BUILD, "coding_build_package")
    summary = build.build(ROOT)
    assert summary["n_sampled"] == 120
    assert summary["shortfalls"] == []
    return build, summary


def test_strata_sizes_and_no_unit_duplicates(built_package) -> None:
    _build, _summary = built_package
    key = pd.read_csv(OUT / "sampling_key.csv")
    assert len(key) == 120
    expected = {
        "COR2": 20,
        "SELF1": 20,
        "HON4": 15,
        "CARE2": 15,
        "AGENT1": 15,
        "OTHER_COR": 15,
        "NONERODE_COR2": 4,
        "NONERODE_SELF1": 4,
        "NONERODE_HON4": 4,
        "NONERODE_CARE2": 4,
        "NONERODE_AGENT1": 4,
    }
    counts = key["stratum"].value_counts().to_dict()
    for k, n in expected.items():
        assert counts.get(k, 0) == n, (k, counts)
    # Unique (config, chain, item) via transition_id: config|...|chain_k|...|item|revise
    units = []
    for tid, item in zip(key["transition_id"], key["item_id"]):
        parts = str(tid).split("|")
        units.append((parts[0], parts[3], str(item)))
    assert len(units) == len(set(units))


def _xlsx_cell_fingerprint(path: Path) -> str:
    from openpyxl import load_workbook
    import hashlib

    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            rows.append(tuple("" if v is None else str(v) for v in row))
    return hashlib.sha256(repr(rows).encode()).hexdigest()


def test_determinism_of_key_and_sheets(built_package, tmp_path: Path) -> None:
    build, _ = built_package
    snap = tmp_path / "snap"
    snap.mkdir()
    for name in ["sampling_key.csv", "sheet_map.csv"]:
        shutil.copy2(OUT / name, snap / name)
    fp_a = _xlsx_cell_fingerprint(OUT / "coderA/coding_sheet_coderA.xlsx")
    fp_b = _xlsx_cell_fingerprint(OUT / "coderB/coding_sheet_coderB.xlsx")

    build.build(ROOT)
    assert filecmp.cmp(snap / "sampling_key.csv", OUT / "sampling_key.csv", shallow=False)
    assert filecmp.cmp(snap / "sheet_map.csv", OUT / "sheet_map.csv", shallow=False)
    assert fp_a == _xlsx_cell_fingerprint(OUT / "coderA/coding_sheet_coderA.xlsx")
    assert fp_b == _xlsx_cell_fingerprint(OUT / "coderB/coding_sheet_coderB.xlsx")


def test_blinding_check_passes(built_package) -> None:
    build, summary = built_package
    assert summary["blinding"]["package_coderA.zip"] is True
    assert summary["blinding"]["package_coderB.zip"] is True
    for zname in ("package_coderA.zip", "package_coderB.zip"):
        result = build.blinding_check_zip(OUT / zname)
        assert result["passed"], result["hits"]
        with zipfile.ZipFile(OUT / zname) as zf:
            names = set(zf.namelist())
            assert "Coding_Guide.pdf" in names
            assert "practice_sheet.xlsx" in names
            assert "Practice_Answers.pdf" in names


def test_analysis_runs_on_synthetic_sheets(built_package) -> None:
    analyze = _load(ANALYZE, "coding_analyze")
    info = analyze.synthetic_dry_run(OUT)
    assert info["n"] == 120
    syn = OUT / "synthetic_results"
    for name in [
        "reliability.json",
        "human_vs_judge.json",
        "themes_by_stratum.csv",
        "q3_by_stratum.csv",
        "disagreements.csv",
        "codes_anonymised.csv",
        "summary.md",
    ]:
        assert (syn / name).exists(), name
    body = (syn / "summary.md").read_text(encoding="utf-8")
    assert "EXPLORATORY (not preregistered)" in body
